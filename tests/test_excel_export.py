import openpyxl

from pipeline.excel_export import export
from pipeline.jobs import insert_job
from pipeline.statuses import create_application


def test_export_writes_four_sheets_with_data(conn, tmp_path):
    jid, _ = insert_job(conn, {"company": "Globex", "title": "Senior Sales Assistant",
                               "source": "wtj", "location": "Lausanne",
                               "discovered_date": "2026-06-11"})
    create_application(conn, jid, source="test", status="Ready to apply")
    out = tmp_path / "out.xlsx"
    export(conn, out, today="2026-06-11")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Tracker", "Summary", "Legend", "New today"]
    tracker = list(wb["Tracker"].iter_rows(values_only=True))
    assert tracker[0][:4] == ("Discovered", "Company", "Role", "Status")
    assert tracker[1][1] == "Globex"
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True, min_row=2)}
    assert summary["Ready to apply"] == 1


def test_export_falls_back_when_target_locked(conn, tmp_path, monkeypatch):
    import pipeline.excel_export as mod
    target = tmp_path / "tracker.xlsx"

    real_replace = mod.os.replace
    def deny(src, dst):
        if str(dst) == str(target):
            raise PermissionError("file is open")
        return real_replace(src, dst)

    monkeypatch.setattr(mod.os, "replace", deny)
    written = export(conn, target, today="2026-06-11")
    assert written != target
    assert written.exists()
    assert "2026-06-11" in written.name
