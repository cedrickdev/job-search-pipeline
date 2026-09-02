"""Regenerate the human-readable Excel report from the DB. Output-only artifact.

Usage: .venv/bin/python -m pipeline.excel_export
"""
import os
import sqlite3
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from pipeline import paths
from pipeline.db import connect, init_db
from pipeline.statuses import STATUS_LEGEND

TRACKER_HEADER = [
    "Discovered", "Company", "Role", "Status", "Score", "Source",
    "Location", "Contract", "Salary",
    "Scored at", "Tailored at", "Approved at", "Applied at",
    "Phone Screen at", "Interview at", "Offer at", "Rejected at",
    "Last event", "Job URL",
]

TRACKER_QUERY = """
SELECT j.discovered_date, j.company, j.title, a.status,
       (SELECT s.score FROM scores s WHERE s.job_id = j.id
        ORDER BY s.created_at DESC LIMIT 1) AS score,
       j.source, j.location, j.contract_type, j.salary,
       a.scored_at, a.tailored_at, a.approved_at, a.submitted_at,
       a.phone_screen_at, a.interview_at, a.offer_at, a.rejected_at,
       (SELECT e.created_at FROM events e WHERE e.application_id = a.id
        ORDER BY e.created_at DESC LIMIT 1) AS last_event,
       j.url
FROM applications a JOIN jobs j ON j.id = a.job_id
ORDER BY j.discovered_date DESC
"""

NEW_TODAY_HEADER = ["Company", "Role", "Score", "Source", "Location", "Job URL"]

NEW_TODAY_QUERY = """
SELECT j.company, j.title,
       (SELECT s.score FROM scores s WHERE s.job_id = j.id
        ORDER BY s.created_at DESC LIMIT 1) AS score,
       j.source, j.location, j.url
FROM jobs j WHERE j.discovered_date = ?
ORDER BY score DESC
"""


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    return ws


def export(conn: sqlite3.Connection, target: str | Path,
           today: str | None = None) -> Path:
    """Write the report atomically. Returns the path actually written
    (a dated fallback if the target is locked by LibreOffice)."""
    target = Path(target)
    today = today or date.today().isoformat()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _sheet(wb, "Tracker", TRACKER_HEADER,
           (tuple(r) for r in conn.execute(TRACKER_QUERY)))
    counts = conn.execute(
        "SELECT status, COUNT(*) FROM applications GROUP BY status ORDER BY COUNT(*) DESC"
    ).fetchall()
    _sheet(wb, "Summary", ["Status", "Count"], (tuple(r) for r in counts))
    _sheet(wb, "Legend", ["Status", "Meaning"], STATUS_LEGEND.items())
    _sheet(wb, "New today", NEW_TODAY_HEADER,
           (tuple(r) for r in conn.execute(NEW_TODAY_QUERY, (today,))))

    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, target)
        return target
    except PermissionError:
        fallback = target.with_name(f"{target.stem}_{today}{target.suffix}")
        os.replace(tmp, fallback)
        return fallback


if __name__ == "__main__":
    conn = connect(paths.DB_PATH)
    init_db(conn)
    written = export(conn, paths.EXPORT_PATH)
    print(f"wrote {written}")
